"""转换核心：Word -> PDF (LibreOffice 无头) 与 PDF -> Word。

PDF -> Word 主引擎用 exactdoc（直接产出高质量可编辑 docx，带 LibreOffice 渲染
闭环校验修正版面偏移）；扫描件/图片型 PDF 先经 ocrmypdf 前置 OCR 叠文字层；
pdf2docx 作为兜底引擎。矢量示意图区域在转换前从 PDF 源头 redact 掉、转换后
渲染成高清图片按锚点插回 docx（复杂排版保真，见 _prepare_figures）。

所有函数返回 (输出路径, 状态) 元组，状态为 "done" / "skipped" / "error"。
冲突处理 conflict: "skip" 跳过已存在, "force" 覆盖。
"""

import os
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

import fitz  # pymupdf

from logging_config import get_logger
logger = get_logger("converters")

# ====== pdf2docx 图片色彩空间补丁（兜底引擎用）======
_pdf2docx_patched = False


def _patch_pdf2docx_images():
    """给 pdf2docx 的图片提取器打补丁，兼容 CMYK / JPX(JPEG2000) 等特殊色彩空间图片。
    只打一次。遇到该错误先转 RGB(A) 再导出 PNG。"""
    global _pdf2docx_patched
    if _pdf2docx_patched:
        return
    try:
        import pymupdf
        import cv2
        import numpy as np
        from pdf2docx.image.ImagesExtractor import ImagesExtractor

        _orig = ImagesExtractor._pixmap_to_cv_image

        def _safe_pixmap_to_cv_image(pixmap):
            try:
                return _orig(pixmap)
            except ValueError:
                for cs in (pymupdf.csRGB, pymupdf.csGRAY):
                    try:
                        rgb = pixmap.convert_to(cs)
                    except Exception:
                        try:
                            rgb = pymupdf.Pixmap(cs, pixmap)
                        except Exception:
                            continue
                    img_byte = rgb.tobytes("png")
                    return cv2.imdecode(np.frombuffer(img_byte, np.uint8),
                                        cv2.IMREAD_UNCHANGED)
                raise

        ImagesExtractor._pixmap_to_cv_image = staticmethod(_safe_pixmap_to_cv_image)
        _pdf2docx_patched = True
        logger.info("已对 pdf2docx 图片导出函数打补丁（兼容 CMYK/JPX 色彩空间）")
    except Exception as e:
        logger.warning("pdf2docx 图片补丁安装失败，特殊色彩空间图片可能丢失: %s", e)


# ====== exactdoc CJK 页码补丁（主引擎用）======
_exactdoc_patched = False


def _patch_exactdoc():
    """exactdoc 0.3.0b1 的 _page_number_gutter 用 int(t) 解析页码装订线，但带圈数字
    (②等 Unicode 数字) isdigit()=True 而 int() 抛 ValueError，导致中文/日文 PDF
    （页脚常用带圈数字页码）主转换崩溃。替换为跳过非标准数字的版本（根因修复）。"""
    global _exactdoc_patched
    if _exactdoc_patched:
        return
    try:
        import exactdoc.infer as m

        def _safe_page_number_gutter(p):
            cands = []
            for bi, blk in enumerate(p.blocks):
                for ln in blk.lines:
                    t = ln.text.strip()
                    if 1 <= len(t) <= 3 and t.isdigit() and ln.horizontal:
                        try:
                            cands.append((bi, ln, int(t)))
                        except ValueError:
                            continue
            if len(cands) < m.GUTTER_MIN_ROWS:
                return []
            cands.sort(key=lambda c: c[1].bbox[2])
            best, cur = [], [cands[0]]
            for c in cands[1:]:
                if c[1].bbox[2] - cur[-1][1].bbox[2] <= m.GUTTER_X_TOL:
                    cur.append(c)
                else:
                    best = max(best, cur, key=len)
                    cur = [c]
            best = max(best, cur, key=len)
            if len(best) < m.GUTTER_MIN_ROWS:
                return []
            best.sort(key=lambda c: c[1].baseline)
            seq = [v for _, _, v in best]
            if sum(1 for a, b in zip(seq, seq[1:]) if b == a + 1) < \
                    m.GUTTER_RUN_SHARE * (len(seq) - 1):
                return []
            if seq[0] != 1:
                return []
            x1 = max(c[1].bbox[2] for c in best)
            ids = {id(c[1]) for c in best}
            if any(ln.bbox[0] < x1 + m.GUTTER_CLEAR_PT
                   for blk in p.blocks for ln in blk.lines if id(ln) not in ids):
                return []
            return best

        m._page_number_gutter = _safe_page_number_gutter
        _exactdoc_patched = True
        logger.info("已对 exactdoc 页码函数打补丁（兼容带圈数字页码，修复 CJK 崩溃）")
    except Exception as e:
        logger.warning("exactdoc 页码补丁安装失败: %s", e)


def find_libreoffice():
    """探测本机 LibreOffice 可执行文件，找不到返回 None。"""
    for name in ("soffice", "libreoffice"):
        exe = shutil.which(name)
        if exe:
            logger.info("找到 LibreOffice(PATH): %s", exe)
            return exe
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice",
        "/snap/bin/libreoffice",
    ]
    for c in candidates:
        if os.path.exists(c):
            logger.info("找到 LibreOffice(常见路径): %s", c)
            return c
    logger.warning("未找到 LibreOffice（PATH 与常见安装路径均未命中）")
    return None


def word_to_pdf(input_path, output_dir, conflict="skip"):
    """Word(.docx/.doc) -> PDF，走 LibreOffice 无头。

    用 -env:UserInstallation 指向临时 profile，避免并发时锁文件冲突。
    """
    soffice = find_libreoffice()
    if not soffice:
        raise RuntimeError("未找到 LibreOffice，请先安装并把 soffice 加入 PATH")
    logger.info("Word→PDF 开始: %s -> %s (conflict=%s)", input_path, output_dir, conflict)
    t0 = time.time()

    src = Path(input_path)
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    target = out / (src.stem + ".pdf")

    if target.exists() and conflict == "skip":
        logger.info("Word→PDF 跳过（已存在且 conflict=skip）: %s", target)
        return str(target), "skipped"
    if target.exists() and conflict == "force":
        target.unlink()

    # 隔离 profile 放到系统临时目录，避免污染输出目录。
    # 注意 file URL 必须是 file:///C:/...（三个斜杠），两斜杠会被 LibreOffice
    # 当成非法路径、profile 起不来，进而报 “bootstrap.ini 已经损坏”。
    profile = Path(tempfile.gettempdir()) / "docflow_lo_profile"
    user_install = "file:///" + str(profile).replace("\\", "/")
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--nofirststartwizard",
        f"-env:UserInstallation={user_install}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(out),
        str(src),
    ]
    logger.debug("LibreOffice cmd: %s", " ".join(cmd))
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0 or not target.exists():
        logger.error("LibreOffice 转换失败 rc=%s stderr=%s",
                     res.returncode, (res.stderr or "").strip())
        raise RuntimeError((res.stderr or res.stdout or "LibreOffice 转换失败").strip())
    logger.info("Word→PDF 完成: %s (耗时 %.2fs)", target, time.time() - t0)
    return str(target), "done"


def _pdf_has_text_layer(path):
    """粗略判断 PDF 是否有文本层（扫描件/图片型返回 False）。

    取样前若干页，平均每页可提取字符 < 5 视为扫描件。检测失败则按有文本层处理，
    交由下游转换引擎报错，避免误判把数字 PDF 送去 OCR 徒增耗时。
    """
    try:
        import fitz  # pymupdf
        doc = fitz.open(path)
        n = len(doc)
        sample = min(n, 12)
        chars = 0
        for i in range(sample):
            chars += len(doc[i].get_text().strip())
        doc.close()
        avg = chars / max(sample, 1)
        has = avg >= 5
        logger.info("文本层检测: 取样 %d/%d 页, 平均 %.1f 字符/页 -> %s",
                    sample, n, avg, "有文本层" if has else "疑似扫描件(无文本层)")
        return has
    except Exception as e:
        logger.warning("文本层检测失败，按有文本层处理: %s", e)
        return True


def _ocr_preprocess(src_pdf):
    """扫描件前置 OCR：用 ocrmypdf + Tesseract 生成带文本层 PDF。

    返回新文件路径（系统临时目录）。失败抛出带安装指引的友好错误。
    """
    try:
        import ocrmypdf
    except Exception:
        raise RuntimeError(
            "该 PDF 为扫描件/图片型（无文字层），需 OCR 才能转成可编辑 Word。"
            "请先安装 OCR 链路：pip install ocrmypdf，并安装系统软件 Tesseract "
            "与其中文语言包 chi_sim（见 README）。")
    tes = _find_tesseract()
    if not tes:
        raise RuntimeError(
            "未检测到 Tesseract OCR 引擎，无法处理扫描件。请安装 Tesseract "
            "（winget install UB-Mannheim.TesseractOCR）并把安装目录加入 PATH。")
    # 保证 ocrmypdf 能在 PATH 里找到 tesseract（安装器有时不自动加 PATH）
    tes_dir = os.path.dirname(tes)
    if tes_dir not in os.environ.get("PATH", ""):
        os.environ["PATH"] = tes_dir + os.pathsep + os.environ.get("PATH", "")
    fd, ocr_out = tempfile.mkstemp(suffix=".pdf", dir=tempfile.gettempdir())
    os.close(fd)
    logger.info("OCR 前置开始: %s -> %s (语言 chi_sim+eng)", src_pdf, ocr_out)
    t0 = time.time()
    try:
        ocrmypdf.ocr(src_pdf, ocr_out, language="chi_sim+eng", force_ocr=True,
                     skip_text=True, output_type="pdf", progress_bar=False)
    except Exception as e:
        try:
            os.remove(ocr_out)
        except Exception:
            pass
        raise RuntimeError(
            f"OCR 处理失败：{e}（请确认 Tesseract 已安装，且中文包 "
            f"chi_sim.traineddata 已放入 Tesseract 的 tessdata 目录）")
    logger.info("OCR 前置完成: 耗时 %.2fs", time.time() - t0)
    return ocr_out


# ---------- 示意图→图片（复杂排版保真，复用 pdf2docx-gui/FreeP2W 验证过的技术） ----------
#
# 原理：矢量画的示意图（框+箭头+标签文字）交给"整页流化"引擎必然错位/乱码
# （内嵌字体缺 ToUnicode 时文字层本身就是错的）。成熟做法是：先在 PDF 源头把
# 示意图区域 redact 掉（删文字+删线稿），引擎只见到干净正文；再把该区域原样
# 渲染成高清图片，按「图后第一段正文」锚点插回 docx 对应位置。图零乱码零偏移，
# 正文/表格/标题照常可编辑。零新增依赖、零模型下载（PyMuPDF 原生能力）。

_TESSERACT_FALLBACK = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def _find_tesseract():
    """定位 tesseract.exe：先 PATH，再常见安装路径。"""
    p = shutil.which("tesseract")
    if p:
        return p
    if os.path.exists(_TESSERACT_FALLBACK):
        return _TESSERACT_FALLBACK
    return None


_MOJI_OK_RANGES = (
    (0x20, 0x7E),      # ASCII
    (0xA0, 0xFF),      # latin-1（· × ° 等）
    (0x0370, 0x03FF),  # 希腊字母（α β γ 数学/科学用法）
    (0x2010, 0x2027),  # 破折号/引号/省略号
    (0x2100, 0x214F),  # ℃ № ™ 等字母符号
    (0x2190, 0x22FF),  # 箭头/数学符号（→ ≤ ∞）
    (0x2460, 0x24FF),  # ① ② ③ 圈码
    (0x2500, 0x25FF),  # 制表符/几何图形
    (0x3000, 0x30FF),  # CJK 标点 + 假名
    (0x3400, 0x9FFF),  # CJK 统一表意
    (0xF900, 0xFAFF),  # CJK 兼容
    (0xFF00, 0xFFEF),  # 全角字符
)


def _is_mojibake(t):
    """文字层乱码判定：异常区块字符占比 ≥35%。

    内嵌字体缺 ToUnicode CMap 时文字层 Unicode 本身就是错的（HTTP→)551、
    TCP→5$1），会掉出正常字符集（泰米尔/格鲁吉亚/叙利亚等区块）。正常
    中英文行几乎不含这些字符，误判概率极低。
    """
    if len(t) < 2:
        return False
    bad = total = 0
    for ch in t:
        if ch.isspace():
            continue
        total += 1
        o = ord(ch)
        if not any(lo <= o <= hi for lo, hi in _MOJI_OK_RANGES):
            bad += 1
    return total > 0 and bad / total >= 0.35


def _page_text_lines(page):
    """取页面正文文字行 [(归一化文本, Rect)]，排除页眉/页脚带与左侧页码栏。"""
    w, h = page.rect.width, page.rect.height
    out = []
    for b in page.get_text("dict")["blocks"]:
        if b.get("type") != 0:
            continue
        for l in b["lines"]:
            t = "".join(s["text"] for s in l["spans"]).strip()
            if not t:
                continue
            r = fitz.Rect(l["bbox"])
            if r.y1 < h * 0.115 or r.y0 > h * 0.91:   # 页眉/页脚（含页首装饰行）
                continue
            if r.x1 < w * 0.12:                       # 侧边页码栏
                continue
            t = "".join(t.split())
            t = "".join(ch for ch in t if ch.isprintable())   # 剔除乱码控制符
            out.append((t, r))
    return out


_CLUSTER_GAP = 24.0   # 证据 y 向间隙小于该值并入同一块


def _detect_figure_regions(pdf):
    """按页检测示意图区域（一页可多个 band）。

    证据 = 矢量路径（表格剔除后）+ 与路径关联的文字行（重叠或紧贴）+ 位图块。
    证据按 y 向相邻度聚类成块：中间隔着「长正文行」（≥20 字且未与路径
    关联的说明段落）的块不合并——这样页首装饰、页中多张图、图间正文
    各自独立，不会连成一个大矩形吞掉正文。每块需 ≥2 条路径，且含 ≥2 行
    图内文字或 ≥1 块位图，才认定为示意图 band；band = 块外包矩形，再吸
    收紧贴的短标签（≤20 字，如图号说明）。渲染成图永远视觉保真，误判
    的代价只是"少一段可编辑文字"，不会出错内容。

    排除项：页眉/页脚装饰与文字、侧边页码栏、表格（PyMuPDF find_tables
    识别，表格保持流化可编辑）。
    """
    regions = []
    for pno in range(pdf.page_count):
        page = pdf[pno]
        w, h = page.rect.width, page.rect.height
        page_area = w * h
        paths = []
        for d in page.get_drawings():
            r = fitz.Rect(d["rect"])
            if r.is_empty or r.get_area() > page_area * 0.85:   # 背景/水印
                continue
            if r.y1 < h * 0.115 or r.y0 > h * 0.91:             # 页眉装饰/页脚
                continue
            if r.x1 < w * 0.12:                                  # 侧边页码栏
                continue
            paths.append(r)
        if len(paths) < 2:      # 单矩形=文本框/装饰框/代码块底框，不当代号图
            continue
        # 表格路径剔除（表格由引擎正常流化，保持可编辑）
        try:
            trs = [fitz.Rect(t.bbox) for t in page.find_tables().tables]
        except Exception:
            trs = []
        if trs:
            def in_table(r):
                for tr in trs:
                    inter = r & tr
                    if not inter.is_empty and \
                            inter.get_area() / max(r.get_area(), 1e-6) >= 0.55:
                        return True
                return False
            paths = [r for r in paths if not in_table(r)]
            if len(paths) < 2:
                continue
        lines = _page_text_lines(page)
        # 与路径关联的文字 = 图内文字（重叠 ≥30% 或中心紧贴路径）
        fig_flags = [False] * len(lines)
        fig_lines = []
        for i, (t, r) in enumerate(lines):
            for p in paths:
                inter = r & p
                if not inter.is_empty and \
                        inter.get_area() / max(r.get_area(), 1e-6) >= 0.3:
                    fig_flags[i] = True
                    fig_lines.append((t, r))
                    break
                if (p + (-8, -8, 8, 8)).contains(
                        fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)):
                    fig_flags[i] = True
                    fig_lines.append((t, r))
                    break
        if len(fig_lines) < 2:
            continue
        # 长正文行 = 聚类障碍：图与图之间的说明段落必须保持可编辑
        barriers = [r for i, (t, r) in enumerate(lines)
                    if len(t) > 20 and not fig_flags[i]]
        # 光栅图片块也是证据：示意图常由「矢量框线 + 位图插画 + 标签」组成，
        # 位图能把同一张图的各部分桥接成一块（整页扫描图除外，那不归我们管）
        ras = []
        for b in page.get_text("dict")["blocks"]:
            if b.get("type") != 1:
                continue
            r = fitz.Rect(b["bbox"])
            if r.is_empty or r.get_area() > page_area * 0.6:
                continue
            if r.y1 < h * 0.115 or r.y0 > h * 0.91:
                continue
            ras.append(r)

        def barrier_between(y0, y1):
            return any(b.y0 > y0 - 2 and b.y1 < y1 + 2 for b in barriers)

        # 竖向聚类：路径 + 图内文字 + 位图块按 y 相邻度并块
        items = ([(r, 1, 0, 0) for r in paths] +
                 [(r, 0, 1, 0) for _, r in fig_lines] +
                 [(r, 0, 0, 1) for r in ras])
        items.sort(key=lambda x: (x[0].y0, x[0].y1))
        clusters = []   # [外包矩形, 路径数, 图内文字行数, 位图块数]
        for r, npath, nfig, nras in items:
            hit = None
            for c in clusters:
                if r.y0 <= c[0].y1 + _CLUSTER_GAP:
                    if r.y0 > c[0].y1 + 2 and barrier_between(c[0].y1, r.y0):
                        continue        # 中间隔着正文段落，不并块
                    hit = c
                    break
            if hit is None:
                clusters.append([fitz.Rect(r), npath, nfig, nras])
            else:
                hit[0] |= r
                hit[1] += npath
                hit[2] += nfig
                hit[3] += nras
        for band, npath, nfig, nras in clusters:
            if npath < 2 or (nfig < 2 and nras < 1):
                continue    # 纯装饰块 / 孤立文本框不是示意图
            band = band & page.rect
            # 吸收紧贴 band 的短文本（示意图侧边/上下标签，如"应用层""图1-3 …"）。
            # 乱码标签（缺 ToUnicode 的图内文字）给更大的窗口——它们绝不是正文，
            # 漏吸就会掉进 Word 正文变成乱码碎片。
            for _ in range(2):
                grown = fitz.Rect(band)
                for t, r in lines:
                    moji = _is_mojibake(t)
                    if len(t) > 20 and not moji:
                        continue
                    cy = (r.y0 + r.y1) / 2
                    up = 45 if moji else 15
                    if not (band.y0 - up <= cy <= band.y1 + 35):
                        continue
                    if r.x0 > band.x1 + 30 or band.x0 > r.x1 + 30:
                        continue
                    grown |= r
                band = grown & page.rect
            if band.get_area() > page_area * 0.9:
                continue
            regions.append((pno, band))
    return regions


def _figure_anchor(pdf, pno, band, bands_by_page):
    """为示意图带选插入锚点：返回 (上方正文, 下方正文) 归一化文本。

    上方=band 之上最后一段正文（定位用），下方=band 之下第一段正文
    （图片插在它前面）。同页没有就顺延到下一页。找不到返回 None。
    """
    def lines_outside_bands(page):
        bands = bands_by_page.get(page.number, [])
        return [(t, r) for t, r in _page_text_lines(page)
                if not any(r.intersects(b) for b in bands)]

    above = below = None
    for t, r in sorted(lines_outside_bands(pdf[pno]), key=lambda x: x[1].y0):
        if len(t) < 6 or _is_mojibake(t):
            continue
        if r.y1 <= band.y0 + 2:
            above = t
        elif r.y0 >= band.y1 - 2 and below is None:
            below = t
            break
    if above is None:       # 图在页首：上方正文顺延到前页找
        for p in (pno - 1, pno - 2):
            if p < 0:
                break
            for t, r in sorted(lines_outside_bands(pdf[p]),
                               key=lambda x: -x[1].y0):
                if len(t) >= 6 and not _is_mojibake(t):
                    above = t
                    break
            if above:
                break
    if below is None:
        for p in (pno + 1, pno + 2):
            if p >= pdf.page_count:
                break
            for t, r in sorted(lines_outside_bands(pdf[p]), key=lambda x: x[1].y0):
                if len(t) >= 6 and not _is_mojibake(t):
                    below = t
                    break
            if below:
                break
    return above, below


def _collect_chrome(pdf):
    """收集每页「页面杂物」的 redact 区域：页眉带、页脚带（含推广水印）、
    侧边原书页码、乱码文字行。

    这些留在 Word 里只能是噪音：原书页码会变成散落正文的孤立数字（甚至
    独占一页），水印每页一行，乱码行是缺 ToUnicode 字体的错误 Unicode，
    任何引擎按文字层转换都必然是错的——从源头删掉才是唯一正确处理。
    """
    chrome = {}
    for pno in range(pdf.page_count):
        page = pdf[pno]
        w, h = page.rect.width, page.rect.height
        rects = [fitz.Rect(0, 0, w, h * 0.115),
                 fitz.Rect(0, h * 0.91, w, h)]
        for b in page.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            for l in b["lines"]:
                t = "".join(s["text"] for s in l["spans"]).strip()
                if not t:
                    continue
                r = fitz.Rect(l["bbox"])
                in_margin = r.x1 < w * 0.12 or r.x0 > w * 0.88
                if (in_margin and t.isdigit()) or _is_mojibake(t):
                    rects.append(r + (-1, -1, 1, 1))
        if rects:
            chrome[pno] = rects
    return chrome


def _prepare_figures(pdf_path):
    """检测示意图 → 渲染 PNG + 选锚点 → 生成 redact 后的工作副本。

    返回 (figures, redacted_path)。figures 每项含 png 字节、宽度(英寸)、
    上/下锚点文本；无示意图返回 ([], None)。任何失败都降级为"无示意图"，
    绝不影响转换主流程。
    """
    try:
        pdf = fitz.open(pdf_path)
    except Exception as e:
        logger.warning("示意图检测跳过（PDF 打不开）: %s", e)
        return [], None
    redacted_path = None
    try:
        regions = _detect_figure_regions(pdf)
        chrome = _collect_chrome(pdf)
        if not regions and not chrome:
            pdf.close()
            return [], None
        logger.info("检测到示意图区域 %d 个、页面杂物 %d 页，开始渲染图片并生成 redact 工作副本",
                    len(regions), len(chrome))
        page_regions = {}
        for pno, rect in regions:
            page_regions.setdefault(pno, []).append(rect)
        figures = []
        ok_regions = []     # 只有确定能插图的区域才允许 redact（绝不丢内容）
        for pno, rect in regions:
            page = pdf[pno]
            clip = (rect + (-4, -4, 4, 4)) & page.rect
            pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), clip=clip, alpha=False)
            above, below = _figure_anchor(pdf, pno, rect, page_regions)
            if not below:
                logger.warning("示意图(p%d)找不到下方锚点，该区域不 redact 不插图", pno + 1)
                continue
            figures.append({
                "png": pix.tobytes("png"),
                "width_in": min(rect.width / 72.0, 6.2),
                "above": above,
                "below": below,
                "page": pno,
            })
            ok_regions.append((pno, rect))
        # redact：从工作副本中删除示意图与页面杂物（图片渲染用的是 redact 前的画面）
        for pno, rect in ok_regions:
            pdf[pno].add_redact_annot(rect + (-2, -2, 2, 2))
        for pno, rects in chrome.items():
            for r in rects:
                pdf[pno].add_redact_annot(r)
        graphics = getattr(fitz, "PDF_REDACT_LINE_ART_REMOVE_IF_COVERED", 1)
        img_mode = getattr(fitz, "PDF_REDACT_IMAGE_REMOVE", 1)   # 位图一并删，
        pages = sorted({p for p, _ in ok_regions} | set(chrome))
        for pno in pages:                                        # 否则位图既被
            pdf[pno].apply_redactions(images=img_mode,           # 单独抽出又
                                      graphics=graphics)         # 在整图里重复
        fd, redacted_path = tempfile.mkstemp(suffix=".pdf", dir=tempfile.gettempdir())
        os.close(fd)
        pdf.save(redacted_path, garbage=3, deflate=True)
        pdf.close()
        logger.info("示意图处理完成: %d 张图待插入, redact 副本 %s", len(figures), redacted_path)
        return figures, redacted_path
    except Exception as e:
        logger.warning("示意图处理失败（按无示意图继续）: %s", e)
        if redacted_path and os.path.exists(redacted_path):
            try:
                os.remove(redacted_path)
            except Exception:
                pass
        redacted_path = None
        return [], None
    finally:
        try:
            pdf.close()
        except Exception:
            pass


def _insert_figure_images(docx_path, figures):
    """把示意图图片按锚点插入 docx。

    锚点匹配范围是全部段落（正文 + 表格单元格，exactdoc 常把图旁文字排进
    表格）；命中的段落在表格里时，图片插到整张表格之前而不是单元格内。
    「下方锚点」找不到时退化为「上方正文段落之后」，再不行才跳过并告警。
    用前向光标保证多张图按文档顺序插入。
    """
    from io import BytesIO
    from docx import Document
    from docx.shared import Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.text.paragraph import Paragraph
    try:
        d = Document(docx_path)
    except Exception as e:
        logger.warning("示意图插入跳过（docx 打不开）: %s", e)
        return
    paras = [Paragraph(p, d) for p in d.element.body.iter(qn("w:p"))]
    norms = ["".join(ch for ch in "".join(p.text.split()) if ch.isprintable())
             for p in paras]

    def new_image_par(ref):
        """在参照段落前建新段落；参照在表格内时插到外层整张表格之前。"""
        el = ref._p
        parent = el.getparent()
        if parent is not None and parent.tag in (qn("w:tc"), qn("w:tbl")):
            tbl = el
            while tbl.getparent() is not None and \
                    tbl.getparent().tag in (qn("w:tc"), qn("w:tr"), qn("w:tbl")):
                tbl = tbl.getparent()
            newp_el = OxmlElement("w:p")
            tbl.addprevious(newp_el)
        else:
            newp_el = OxmlElement("w:p")
            el.addprevious(newp_el)
        return Paragraph(newp_el, d)

    # 页码→docx 位置校准：成对命中且位置合理的锚点当路标，其余图按 PDF 页码
    # 在路标间插值估计位置。没有位置校验时，正文里的重复语句会把光标越推
    # 越远，后面几十张图的锚点会全部搜空。
    n_pages = max(f["page"] for f in figures) + 1
    par_rate = len(paras) / max(n_pages, 1)
    tol = max(3 * par_rate, 90)
    calib = []          # [(pdf_page, docx_pos)]，按页升序

    def expected_pos(page):
        if not calib:
            return page * par_rate
        prev = next((c for c in reversed(calib) if c[0] <= page), None)
        nxt = next((c for c in calib if c[0] > page), None)
        if prev and nxt and nxt[0] > prev[0]:
            k = (page - prev[0]) / (nxt[0] - prev[0])
            return prev[1] + k * (nxt[1] - prev[1])
        if prev:
            return prev[1] + (page - prev[0]) * par_rate
        return nxt[1] - (nxt[0] - page) * par_rate

    cursor = 0
    inserted = 0
    for idx, fig in enumerate(figures):
        above = fig.get("above")
        below = fig["below"]
        exp = expected_pos(fig["page"])
        hit = -1
        if above:
            # (above, below) 成对匹配 + 位置校验：目录里的同名标题、后文
            # 重复语句都会被「离估计位置太远」拦下
            best = -1
            best_d = None
            for a in [i for i in range(cursor, len(paras))
                      if above in norms[i]][:20]:
                for i in range(a + 1, min(a + 61, len(paras))):
                    if below in norms[i] and abs(i - exp) <= tol:
                        dist = abs(i - exp)
                        if best_d is None or dist < best_d:
                            best, best_d = i, dist
                        break
            if best >= 0:
                hit = best
                calib.append((fig["page"], hit))
                calib.sort(key=lambda c: c[0])
            else:
                a0 = next((i for i in range(cursor, len(paras))
                           if above in norms[i]), -1)
                if a0 >= 0 and abs(a0 + 1 - exp) <= tol:
                    hit = a0 + 1     # 退化：插在首个上方正文之后
        if hit < 0:
            b_pos = [i for i in range(cursor, len(paras)) if below in norms[i]]
            if b_pos:
                c = min(b_pos, key=lambda i: (abs(i - exp), i))
                if abs(c - exp) <= tol:
                    hit = c
        if hit < 0:
            # 兜底：按估计位置直接插——图已从 PDF 删掉，宁可位置略偏不丢图
            hit = max(cursor, min(int(exp), len(paras) - 1))
            logger.warning("示意图[%d]锚点未命中，按页码估计位置插入(p%d)",
                           idx, fig["page"] + 1)
        try:
            if hit >= len(paras):
                newp = d.add_paragraph()
            else:
                newp = new_image_par(paras[hit])
            newp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = newp.add_run()
            run.add_picture(BytesIO(fig["png"]), width=Inches(fig["width_in"]))
            cursor = max(cursor, min(hit, len(paras) - 1))
            inserted += 1
        except Exception as e:
            logger.warning("示意图[%d]插入失败: %s", idx, e)
    if inserted:
        d.save(docx_path)
    logger.info("示意图插入完成: %d/%d 张", inserted, len(figures))


def _pdf_to_word_pdf2docx(input_path, target_path, conflict="skip"):
    """兜底引擎：pdf2docx 主路径（convert 到指定 target 路径）。"""
    from pdf2docx import Converter  # 懒加载
    _patch_pdf2docx_images()
    target = Path(target_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and conflict == "skip":
        return str(target), "skipped"
    if target.exists():
        target.unlink()
    t0 = time.time()
    cv = Converter(str(input_path))
    try:
        cv.convert(str(target))
    finally:
        cv.close()
    if not target.exists():
        raise RuntimeError("pdf2docx 未生成输出文件")
    logger.info("PDF→Word(pdf2docx兜底) 完成: %s (耗时 %.2fs, %d bytes)",
                target, time.time() - t0, target.stat().st_size)
    return str(target), "done"


def _apply_outline_as_headings(pdf_path, docx_path, max_heading=3):
    """把 PDF 大纲（书签/章节树）映射成 Word 标题样式，使生成的 docx 在
    Word 导航窗格里可点击跳转章节。

    - 用 fitz 取大纲 [层级, 标题, 页码]；
    - 用 python-docx 打开 docx，按「归一化文本」匹配段落，套 Heading 1/2/3；
    - 归一化：去掉所有空白（含全角空格），避免 exactdoc 与 PDF 间空格差异；
    - 匹配策略：先在全文找「段落文本 == 标题」的精确匹配，找不到再放宽到
      「段落基本就是该标题（仅多 ≤8 个归一化字符）」；用前向光标保证单调对齐，
      避免正文里提到章节名被误标；
    - 找不到对应正文的条目（如「封面」「版权页」纯书签）直接跳过，绝不影响转换。
    """
    try:
        import fitz
        from docx import Document
    except Exception as e:
        logger.warning("大纲→标题样式跳过（依赖缺失 fitz/python-docx）: %s", e)
        return

    try:
        doc = fitz.open(pdf_path)
        toc = doc.get_toc()
        doc.close()
    except Exception as e:
        logger.warning("读取 PDF 大纲失败，跳过标题样式: %s", e)
        return
    if not toc:
        logger.info("PDF 无大纲，跳过标题样式")
        return

    try:
        d = Document(docx_path)
    except Exception as e:
        logger.warning("打开 docx 失败，跳过标题样式: %s", e)
        return

    def norm(s):
        return "".join(s.split())  # 去掉所有空白（含全角空格）

    paras = d.paragraphs
    norm_paras = [norm(p.text) for p in paras]
    used = [False] * len(paras)
    cursor = 0
    applied = 0
    skipped = 0

    for level, title, _page in toc:
        nt = norm(title)
        if len(nt) < 2:
            skipped += 1
            continue
        hstyle = "Heading %d" % min(level, max_heading)
        hit = -1
        # 1) 精确相等
        for i in range(cursor, len(paras)):
            if not used[i] and norm_paras[i] == nt:
                hit = i
                break
        # 2) 放宽：段落基本就是该标题（仅多 ≤8 个归一化字符）
        if hit < 0:
            for i in range(cursor, len(paras)):
                if (not used[i] and nt in norm_paras[i]
                        and len(norm_paras[i]) <= len(nt) + 8):
                    hit = i
                    break
        if hit < 0:
            skipped += 1
            continue
        try:
            paras[hit].style = d.styles[hstyle]
            used[hit] = True
            cursor = hit + 1
            applied += 1
        except Exception as e:
            logger.debug("应用标题样式失败 [%s]%s: %s", hstyle, title, e)
            skipped += 1

    if applied:
        d.save(docx_path)
        logger.info("大纲→标题样式：应用 %d 个（跳过 %d 个无正文对应），层级上限 Heading %d",
                    applied, skipped, max_heading)
    else:
        logger.info("大纲→标题样式：未找到可匹配章节标题（%d 个跳过），不改动 docx", skipped)


def pdf_to_word(input_path, output_dir, conflict="skip", options=None):
    """PDF -> Word(.docx)。

    主引擎 exactdoc（直接产出高质量可编辑 docx + LibreOffice 渲染闭环校验修正版面偏移）；
    扫描件先经 ocrmypdf 前置 OCR 叠文字层；exactdoc 失败时回退 pdf2docx。
    转换前把矢量示意图区域从 PDF 源头 redact 掉（引擎只见干净正文），转换后把
    该区域渲染成高清图片按锚点插回 docx——图零乱码零偏移（_prepare_figures）。
    转换完成后固定把 PDF 大纲映射成 Word 标题样式（章节导航是默认行为，无开关）。
    """
    src = Path(input_path)
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    target = out / (src.stem + ".docx")
    if target.exists() and conflict == "skip":
        logger.info("PDF→Word 跳过（已存在且 conflict=skip）: %s", target)
        return str(target), "skipped"
    if target.exists() and conflict == "force":
        target.unlink()

    source = str(src)
    work = source
    ocr_tmp = None
    tmp_files = []
    figures = []
    try:
        # 扫描件前置 OCR
        if not _pdf_has_text_layer(str(src)):
            ocr_tmp = _ocr_preprocess(str(src))
            source = ocr_tmp
            work = source

        # 示意图：源头 redact + 渲染成图（无示意图时为空操作）
        figures, redacted = _prepare_figures(source)
        if redacted:
            work = redacted
            tmp_files.append(redacted)

        # 主引擎 exactdoc
        _patch_exactdoc()
        from exactdoc import convert as exact_convert
        logger.info("PDF→Word(exactdoc) 开始: %s -> %s (refine_rounds=2)", work, target)
        t0 = time.time()
        exact_convert(work, str(target), oracle="libreoffice",
                      refine_rounds=2, max_pages=0)
        if not target.exists():
            raise RuntimeError("exactdoc 未生成输出文件")
        logger.info("PDF→Word(exactdoc) 完成: %s (耗时 %.2fs, %d bytes)",
                    target, time.time() - t0, target.stat().st_size)
        _apply_outline_as_headings(str(src), str(target))
        if figures:
            _insert_figure_images(str(target), figures)
        return str(target), "done"
    except Exception as e:
        logger.warning("exactdoc 主路径失败，回退 pdf2docx: %s", e)
        out_path, st = _pdf_to_word_pdf2docx(work, str(target), conflict="force")
        if st == "done":
            _apply_outline_as_headings(str(src), str(target))
            if figures:
                _insert_figure_images(str(target), figures)
        return out_path, st
    finally:
        if ocr_tmp and os.path.exists(ocr_tmp):
            try:
                os.remove(ocr_tmp)
            except Exception:
                pass
        for f in tmp_files:
            try:
                os.remove(f)
            except Exception:
                pass


def convert(tool_id, input_path, output_dir, conflict="skip", options=None):
    """按工具 id 分发到具体转换器。"""
    logger.info("convert 分发: tool_id=%s, input=%s, out=%s, conflict=%s, options=%s",
                tool_id, input_path, output_dir, conflict, options)
    if tool_id == "word2pdf":
        return word_to_pdf(input_path, output_dir, conflict)
    if tool_id == "pdf2word":
        return pdf_to_word(input_path, output_dir, conflict, options)
    logger.error("未实现的工具: %s", tool_id)
    raise ValueError(f"未实现的工具: {tool_id}")
